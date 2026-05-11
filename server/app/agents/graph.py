from __future__ import annotations

import json
import logging
import uuid
from datetime import datetime, timezone

from langgraph.graph import END, START, StateGraph

from app.agents.analyser import build_analyser_subgraph
from app.agents.discovery import build_discovery_subgraph
from app.agents.sprint import build_sprint_subgraph
from app.services.persistence import save_artifact, save_index
from app.services.rag import build_working_index
from app.shared.state_types import GraphState


log = logging.getLogger(__name__)


def _now_iso() -> str:
    """Return current UTC time in ISO format."""
    return datetime.now(timezone.utc).isoformat()


def ingest_node(state: GraphState) -> dict:
    """Prepare minimal parsed document structure before analysis starts."""
    log.info("[ingest_node] START project_id=%s parsed_docs=%d", state.get("project_id"), len(state.get("parsed_documents", [])))
    updates: dict = {}

    if not state.get("parsed_documents"):
        parsed_documents = []
        for file_name in state.get("raw_files", []):
            parsed_documents.append(
                {
                    "file_name": file_name,
                    "file_type": "text/plain",
                    "s3_key": f"uploads/{state['project_id']}/{file_name}",
                    "sections": [
                        {
                            "file_name": file_name,
                            "section_heading": "Imported Context",
                            "page": 1,
                            "content_type": "text",
                            "content": state.get("additional_context", ""),
                            "raw_image_ref": None,
                        }
                    ],
                }
            )
        updates["parsed_documents"] = parsed_documents

    if state.get("review_1_status") is None:
        updates["review_1_status"] = "pending"
    if state.get("review_2_status") is None:
        updates["review_2_status"] = "pending"

    return updates


def raw_rag_index_node(state: GraphState) -> dict:
    """Build a working RAG index from parsed sections."""
    chunk_ids, index_path = build_working_index(
        project_id=state["project_id"],
        version=state["version"],
        parsed_documents=state.get("parsed_documents", []),
    )

    return {
        "working_chunk_ids": chunk_ids,
        "streaming_events": [
            {
                "event_id": str(uuid.uuid4()),
                "type": "working_index_created",
                "node": "raw_rag_index_node",
                "payload": {
                    "index_path": index_path,
                    "chunk_count": len(chunk_ids),
                },
                "timestamp": _now_iso(),
            }
        ],
    }


def apply_review_1_edits_node(state: GraphState) -> dict:
    """Apply review-1 human edits from payload into analyser output."""
    payload = state.get("user_edits_payload") or {}
    analyser = state.get("analyser_output")
    if not analyser:
        return {"review_1_status": "approved"}

    deltas = []
    summary = payload.get("executive_summary")
    if isinstance(summary, str) and summary.strip() and summary != analyser.get("executive_summary"):
        deltas.append(
            {
                "change_id": str(uuid.uuid4()),
                "source": "user_edit",
                "source_ref": "review_1",
                "field_path": "analyser_output.executive_summary",
                "old_value": analyser.get("executive_summary"),
                "new_value": summary.strip(),
                "timestamp": _now_iso(),
            }
        )
        analyser = {**analyser, "executive_summary": summary.strip()}

    return {
        "analyser_output": analyser,
        "review_1_status": "approved",
        "delta_changes": deltas,
    }


def route_review_2_node(state: GraphState) -> dict:
    """No-op node used as a clean interrupt checkpoint before final routing."""
    _ = state
    return {}


def apply_review_2_edits_node(state: GraphState) -> dict:
    """Apply final human markdown edits before artifact export."""
    payload = state.get("user_edits_payload") or {}
    appendix = payload.get("final_doc_appendix")
    markdown = state.get("final_doc_markdown") or ""
    if not isinstance(appendix, str) or not appendix.strip():
        return {"review_2_status": "approved"}

    updated_markdown = f"{markdown}\n\n## Review Notes\n\n{appendix.strip()}\n"
    return {
        "final_doc_markdown": updated_markdown,
        "review_2_status": "approved",
        "delta_changes": [
            {
                "change_id": str(uuid.uuid4()),
                "source": "user_edit",
                "source_ref": "review_2",
                "field_path": "final_doc_markdown",
                "old_value": markdown,
                "new_value": updated_markdown,
                "timestamp": _now_iso(),
            }
        ],
    }


def approved_rag_index_node(state: GraphState) -> dict:
    """Save approved requirements/risks index for retrieval after sign-off."""
    analyser = state.get("analyser_output") or {}
    records = []

    for req in analyser.get("functional_requirements", []):
        records.append(
            {
                "kind": "requirement",
                "id": req.get("req_id"),
                "text": req.get("description", ""),
                "metadata": json.dumps({"moscow": req.get("moscow")}),
            }
        )
    for risk in analyser.get("risks", []):
        records.append(
            {
                "kind": "risk",
                "id": risk.get("risk_id"),
                "text": risk.get("description", ""),
                "metadata": json.dumps({"severity": risk.get("severity")}),
            }
        )

    index_path = save_index(
        project_id=state["project_id"],
        version=state["version"],
        kind="approved",
        records=records,
    )

    return {
        "streaming_events": [
            {
                "event_id": str(uuid.uuid4()),
                "type": "approved_index_created",
                "node": "approved_rag_index_node",
                "payload": {"index_path": index_path, "record_count": len(records)},
                "timestamp": _now_iso(),
            }
        ]
    }


def artifact_export_node(state: GraphState) -> dict:
    """Export markdown into md/pdf/docx artifacts.

    We try real generation first, then fallback to text placeholders.
    """
    markdown = state.get("final_doc_markdown") or "# Business Requirements Analysis\n\n_No content available._\n"
    md_key = save_artifact(state["project_id"], state["version"], "md", markdown)

    pdf_content: bytes | str
    try:
        from weasyprint import HTML

        pdf_content = HTML(string=f"<pre>{markdown}</pre>").write_pdf()
    except Exception:
        pdf_content = f"PDF_PLACEHOLDER\n\nGenerated from:\n{markdown}"

    try:
        from docx import Document
        from io import BytesIO

        document = Document()
        for line in markdown.splitlines():
            document.add_paragraph(line)
        buf = BytesIO()
        document.save(buf)
        docx_content: bytes | str = buf.getvalue()
    except Exception:
        docx_content = f"DOCX_PLACEHOLDER\n\nGenerated from:\n{markdown}"

    pdf_key = save_artifact(state["project_id"], state["version"], "pdf", pdf_content)
    docx_key = save_artifact(state["project_id"], state["version"], "docx", docx_content)

    return {
        "final_doc_markdown": markdown,
        "final_doc_pdf_s3_key": pdf_key,
        "final_doc_docx_s3_key": docx_key,
        "streaming_events": [
            {
                "event_id": str(uuid.uuid4()),
                "type": "artifacts_exported",
                "node": "artifact_export_node",
                "payload": {
                    "markdown": md_key,
                    "pdf": pdf_key,
                    "docx": docx_key,
                },
                "timestamp": _now_iso(),
            }
        ],
    }


def _route_after_analyser(state: GraphState) -> str:
    """Choose next node after Stage-1 human review decision."""
    return "discovery_stage_node" if state["review_1_status"] == "approved" else "apply_review_1_edits_node"


def _route_after_review_2(state: GraphState) -> str:
    """Choose next node after Stage-2 human review decision."""
    if state["review_2_status"] == "more_questions":
        return "discovery_stage_node"
    if state["review_2_status"] == "approved":
        return "approved_rag_index_node"
    return "apply_review_2_edits_node"


def mock_architecture_node(state: GraphState) -> dict:
    """Generate mocked architecture diagrams for Stage 3 (until real Architecture Agent is ready)."""
    import uuid as _uuid
    from datetime import datetime, timezone

    project_id = state.get("project_id", "unknown")
    analyser = state.get("analyser_output") or {}
    team = analyser.get("recommended_team", {})
    team_roles = team.get("roles", ["Business Analyst", "Tech Lead", "QA Engineer"])
    team_size = team.get("size", 3)

    arch_tech_notes = (
        "Backend: FastAPI + Python 3.12. "
        "Agent framework: LangGraph. "
        "Database: SQLite (dev) → PostgreSQL (prod). "
        f"Team roles derived from analysis: {', '.join(team_roles)}."
    )

    diagrams = [
        {
            "diagram_id": str(_uuid.uuid4()),
            "title": "System Architecture",
            "diagram_type": "system_architecture",
            "tool": "plantuml",
            "dsl": (
                "@startuml\n!theme plain\n"
                "package \"Frontend\" {\n  [React SPA] as UI\n}\n"
                "package \"API Layer\" {\n  [FastAPI] as API\n  [WebSocket Handler] as WS\n}\n"
                "package \"Agent Layer\" {\n"
                "  [Analyser Agent] as Analyser\n"
                "  [Discovery Agent] as Discovery\n"
                "  [Architecture Agent] as Architecture\n"
                "  [Sprint Agent] as Sprint\n}\n"
                "package \"Data Layer\" {\n  database \"SQLite\" as DB\n  folder \"File Storage\" as FS\n}\n"
                "UI --> API : REST calls\nUI --> WS : WebSocket stream\n"
                "API --> Analyser : invoke\nAnalyser --> Discovery : state handoff\n"
                "Discovery --> Architecture : state handoff\nArchitecture --> Sprint : state handoff\n"
                "API --> DB : read/write\nAPI --> FS : upload/download\n@enduml\n"
            ),
            "description": "Component-level view of the full BRA Tool system.",
        },
        {
            "diagram_id": str(_uuid.uuid4()),
            "title": "Data Flow",
            "diagram_type": "data_flow",
            "tool": "mermaid",
            "dsl": (
                "flowchart TD\n"
                "    User([User / Product Manager]) -->|Upload Docs| UploadUI[Upload UI]\n"
                "    UploadUI -->|multipart/form-data| FastAPI[FastAPI Backend]\n"
                "    FastAPI -->|parse + chunk| DocParser[Document Parser]\n"
                "    DocParser -->|sections| RAGIndex[(RAG Index)]\n"
                "    FastAPI -->|invoke graph| LangGraph[LangGraph Pipeline]\n"
                "    LangGraph -->|read chunks| RAGIndex\n"
                "    LangGraph -->|write outputs| SQLite[(SQLite DB)]\n"
                "    LangGraph -->|stream events| WebSocket[WebSocket /ws]\n"
                "    WebSocket -->|live tokens| FrontendUI[React Frontend]\n"
                "    SQLite -->|query| FastAPI\n"
                "    FastAPI -->|JSON response| FrontendUI\n"
            ),
            "description": "How data moves from document upload to sprint plan generation.",
        },
        {
            "diagram_id": str(_uuid.uuid4()),
            "title": "User Flow",
            "diagram_type": "user_flow",
            "tool": "mermaid",
            "dsl": (
                "sequenceDiagram\n"
                "    participant PM as Product Manager\n"
                "    participant UI as React UI\n"
                "    participant API as FastAPI\n"
                "    participant Graph as LangGraph\n"
                "    PM->>UI: Upload requirement docs\n"
                "    UI->>API: POST /projects/{id}/files\n"
                "    PM->>UI: Trigger analysis\n"
                "    UI->>API: POST /projects/{id}/run\n"
                "    API->>Graph: Run Analyser + Discovery stages\n"
                "    Graph-->>API: Stream events via WebSocket\n"
                "    API-->>UI: Live status updates\n"
                "    PM->>UI: Review & approve Stage 1\n"
                "    UI->>API: POST /projects/{id}/discovery/answer\n"
                "    PM->>UI: Approve architecture (Stage 3)\n"
                "    UI->>API: POST /projects/{id}/architecture/approve\n"
                "    API->>Graph: Run Sprint Agent (Stage 4)\n"
                "    Graph-->>API: SprintPlan generated\n"
                "    PM->>UI: Review & approve Sprint Plan\n"
                "    UI->>API: POST /projects/{id}/sprint/approve\n"
            ),
            "description": "End-to-end journey for Product Manager through all 4 stages.",
        },
    ]

    architecture_output = {
        "diagrams": diagrams,
        "tech_stack_notes": arch_tech_notes,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "is_mocked": True,
    }

    return {
        "architecture_output": architecture_output,
        "review_3_status": "pending",
        "streaming_events": [
            {
                "event_id": str(_uuid.uuid4()),
                "type": "architecture_generated",
                "node": "mock_architecture_node",
                "payload": {
                    "diagram_count": len(diagrams),
                    "is_mocked": True,
                },
                "timestamp": datetime.now(timezone.utc).isoformat(),
            }
        ],
    }


def apply_review_3_edits_node(state: GraphState) -> dict:
    """Apply Stage-3 architecture approval and mark review_3_status as approved."""
    payload = state.get("user_edits_payload") or {}
    arch = state.get("architecture_output") or {}

    # Apply any user diagram edits if provided
    if "diagrams" in payload and isinstance(payload["diagrams"], list):
        arch = {**arch, "diagrams": payload["diagrams"]}

    return {
        "architecture_output": arch,
        "review_3_status": "approved",
    }


def apply_review_4_edits_node(state: GraphState) -> dict:
    """Apply Stage-4 sprint plan approval and mark review_4_status as approved."""
    payload = state.get("user_edits_payload") or {}
    sprint_plan = state.get("sprint_plan") or {}

    # Append reviewer notes if provided
    notes = payload.get("sprint_notes") or payload.get("notes")
    if notes:
        sprint_plan = {**sprint_plan, "reviewer_notes": notes}

    return {
        "sprint_plan": sprint_plan,
        "review_4_status": "approved",
    }


def build_parent_graph():
    """Build the full parent graph that hosts analyser and discovery subgraphs."""
    analyser_subgraph = build_analyser_subgraph()
    discovery_subgraph = build_discovery_subgraph()

    graph = StateGraph(GraphState)

    graph.add_node("ingest_node", ingest_node)
    graph.add_node("raw_rag_index_node", raw_rag_index_node)
    graph.add_node("analyser_stage_node", analyser_subgraph)
    graph.add_node("apply_review_1_edits_node", apply_review_1_edits_node)
    graph.add_node("discovery_stage_node", discovery_subgraph)
    graph.add_node("route_review_2_node", route_review_2_node)
    graph.add_node("apply_review_2_edits_node", apply_review_2_edits_node)
    graph.add_node("approved_rag_index_node", approved_rag_index_node)
    graph.add_node("artifact_export_node", artifact_export_node)

    graph.add_edge(START, "ingest_node")
    graph.add_edge("ingest_node", "raw_rag_index_node")
    graph.add_edge("raw_rag_index_node", "analyser_stage_node")
    graph.add_conditional_edges(
        "analyser_stage_node",
        _route_after_analyser,
        {
            "apply_review_1_edits_node": "apply_review_1_edits_node",
            "discovery_stage_node": "discovery_stage_node",
        },
    )
    graph.add_edge("apply_review_1_edits_node", "discovery_stage_node")
    graph.add_edge("discovery_stage_node", "route_review_2_node")
    graph.add_conditional_edges(
        "route_review_2_node",
        _route_after_review_2,
        {
            "apply_review_2_edits_node": "apply_review_2_edits_node",
            "discovery_stage_node": "discovery_stage_node",
            "approved_rag_index_node": "approved_rag_index_node",
        },
    )
    graph.add_edge("apply_review_2_edits_node", "approved_rag_index_node")
    graph.add_edge("approved_rag_index_node", "artifact_export_node")
    graph.add_edge("artifact_export_node", END)

    return graph.compile(
        interrupt_before=["apply_review_1_edits_node", "route_review_2_node"]
    )

def get_graph():
    """Return compiled parent graph instance."""
    return build_parent_graph()
